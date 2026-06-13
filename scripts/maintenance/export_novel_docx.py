#!/usr/bin/env python3
"""Merge final novel chapters into a navigable Word document."""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
LOCAL_DEPS = REPO_ROOT / ".artifacts" / "docx_deps"
if LOCAL_DEPS.exists():
    sys.path.insert(0, str(LOCAL_DEPS))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


CHAPTER_RE = re.compile(r"chapter_(\d{4})\.txt$")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def set_run_font(run, name: str, size: float, *, bold: bool | None = None, color: str = "222222") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_field(run, instruction: str) -> None:
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, end])


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_link(paragraph, text: str, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "333333")
    props.append(color)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "宋体")
    fonts.set(qn("w:ascii"), "SimSun")
    fonts.set(qn("w:hAnsi"), "SimSun")
    props.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "21")
    props.append(size)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_markdown_runs(paragraph, text: str) -> None:
    text = text.replace("`", "")
    text = re.sub(r"(?<!\*)\*(?!\*)", "", text)
    cursor = 0
    for match in BOLD_RE.finditer(text):
        if match.start() > cursor:
            set_run_font(paragraph.add_run(text[cursor:match.start()]), "宋体", 11)
        set_run_font(paragraph.add_run(match.group(1)), "宋体", 11, bold=True)
        cursor = match.end()
    if cursor < len(text):
        set_run_font(paragraph.add_run(text[cursor:]), "宋体", 11)


def chapter_title(project: Path, chapter: int) -> str:
    outline = project / "chapters" / "outline" / f"chapter_{chapter:04d}.json"
    title = ""
    if outline.exists():
        try:
            title = str(json.loads(outline.read_text(encoding="utf-8")).get("title", "")).strip()
        except Exception:
            title = ""
    return f"第{chapter}章 {title}" if title else f"第{chapter}章"


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    fmt = normal.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(0.74)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(5)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    heading = doc.styles["Heading 1"]
    heading.font.name = "黑体"
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    heading.font.size = Pt(16)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(46, 74, 92)
    heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(18)
    heading.paragraph_format.keep_with_next = True

    if "Contents Entry" not in [style.name for style in doc.styles]:
        contents = doc.styles.add_style("Contents Entry", WD_STYLE_TYPE.PARAGRAPH)
    else:
        contents = doc.styles["Contents Entry"]
    contents.font.name = "宋体"
    contents._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    contents.font.size = Pt(10.5)
    contents.paragraph_format.left_indent = Cm(0.3)
    contents.paragraph_format.space_after = Pt(2)
    contents.paragraph_format.line_spacing = 1.15


def configure_section(section, title: str, *, first: bool = False) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    section.different_first_page_header_footer = first

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_p.text = title
    if header_p.runs:
        set_run_font(header_p.runs[0], "宋体", 9, color="777777")

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run()
    set_run_font(run, "宋体", 9, color="777777")
    add_field(run, " PAGE ")


def load_chapters(project: Path, selected: list[int] | None = None) -> list[tuple[int, Path, str]]:
    results = []
    for path in sorted((project / "chapters" / "final").glob("chapter_*.txt")):
        match = CHAPTER_RE.match(path.name)
        if not match:
            continue
        chapter = int(match.group(1))
        if selected is not None and chapter not in selected:
            continue
        results.append((chapter, path, chapter_title(project, chapter)))
    return results


def build_docx(project: Path, output: Path, selected: list[int] | None = None) -> None:
    world = json.loads((project / "world.json").read_text(encoding="utf-8"))
    title = str(world.get("title", project.name)).strip()
    subtitle = str(world.get("subtitle", "")).strip()
    chapters = load_chapters(project, selected)
    if not chapters:
        raise RuntimeError("未找到 final 章节")

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], title, first=True)
    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Cm(6.2)
    cover.paragraph_format.space_after = Pt(12)
    set_run_font(cover.add_run(title), "黑体", 30, bold=True, color="172A3A")
    if subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.paragraph_format.space_after = Pt(24)
        set_run_font(sub.add_run(subtitle), "宋体", 15, color="555555")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Cm(4)
    set_run_font(meta.add_run(f"全本合订版 · 共{len(chapters)}章"), "宋体", 11, color="777777")
    doc.add_page_break()

    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.space_after = Pt(16)
    set_run_font(toc_title.add_run("目录"), "黑体", 20, bold=True, color="172A3A")
    for chapter, _, display_title in chapters:
        p = doc.add_paragraph(style="Contents Entry")
        add_internal_link(p, display_title, f"chapter_{chapter:04d}")
    doc.add_page_break()

    bookmark_id = 1
    for index, (chapter, path, display_title) in enumerate(chapters):
        if index:
            doc.add_page_break()
        heading = doc.add_paragraph(display_title, style="Heading 1")
        add_bookmark(heading, f"chapter_{chapter:04d}", bookmark_id)
        bookmark_id += 1

        text = path.read_text(encoding="utf-8", errors="ignore").strip()
        for raw in re.split(r"\n\s*\n", text):
            value = raw.strip()
            if not value:
                continue
            paragraph = doc.add_paragraph()
            add_markdown_runs(paragraph, value.replace("\n", ""))

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = title
    doc.core_properties.subject = subtitle or "长篇小说合订本"
    doc.core_properties.author = "Novel Generation Workflow"
    doc.core_properties.keywords = "小说, 合订本, 目录"
    doc.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="将 final 章节合并为带目录的 Word 小说")
    parser.add_argument("--project", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--chapters", default="", help="抽样章节，如 1,2,500")
    args = parser.parse_args()

    selected = None
    if args.chapters.strip():
        selected = [int(item.strip()) for item in args.chapters.split(",") if item.strip()]
    build_docx(Path(args.project).resolve(), Path(args.output).resolve(), selected)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
